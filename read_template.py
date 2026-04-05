"""Extract structure from the DOCX template to understand formatting."""
from docx import Document
from docx.shared import Pt, Inches

doc = Document("Completed_Paper Template Colloquium 2026.docx")

print("=== DOCUMENT STYLES ===")
for style in doc.styles:
    if style.type is not None and style.type.name == 'PARAGRAPH':
        if style.font.size:
            print(f"  Style: {style.name}, Font: {style.font.name}, Size: {style.font.size}, Bold: {style.font.bold}")

print("\n=== SECTIONS ===")
for i, section in enumerate(doc.sections):
    print(f"  Section {i}: margins L={section.left_margin}, R={section.right_margin}, T={section.top_margin}, B={section.bottom_margin}")
    print(f"    Page: {section.page_width} x {section.page_height}")
    print(f"    Columns: {section._sectPr.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cols')}")

print("\n=== PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    style = p.style.name if p.style else "None"
    text = p.text[:120] if p.text else ""
    alignment = p.alignment
    fmt = ""
    if p.runs:
        r = p.runs[0]
        fmt = f"Font={r.font.name}, Size={r.font.size}, Bold={r.font.bold}, Italic={r.font.italic}"
    print(f"  [{i:3d}] Style={style:30s} Align={str(alignment):10s} {fmt}")
    if text:
        print(f"        TEXT: {text}")
    print()

print("\n=== TABLES ===")
for i, table in enumerate(doc.tables):
    print(f"  Table {i}: {len(table.rows)} rows x {len(table.columns)} cols")
    for j, row in enumerate(table.rows):
        cells = [cell.text[:40] for cell in row.cells]
        print(f"    Row {j}: {cells}")
